"""Editable Office exports and browser print views for user-owned content."""

from __future__ import annotations

import html
import io
import math
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt as DocxPt, RGBColor as DocxRGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from . import documents

MIME = {
    "word": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "ppt": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
EXT = {"word": "docx", "excel": "xlsx", "ppt": "pptx"}

CATALOG = [
    {"id": "rental-contract", "kind": "word", "title": "房屋租赁合同", "description": "可编辑的租赁合同空白框架；与 Agent 对话可生成已填好的版本。",
     "source_url": documents.TEMPLATES["rental-contract"]["source_url"]},
    {"id": "evidence-ledger", "kind": "excel", "title": "案件证据清单", "description": "按证据名称、来源、待证明事实和文件位置整理材料。"},
    {"id": "case-brief", "kind": "ppt", "title": "案件沟通简报", "description": "用于与律师沟通的事实、时间线和待确认问题框架。"},
]


def _font(style, size: int, bold: bool = False, color: str = "243F58") -> None:
    style.font.name = "PingFang SC"
    style.font.size = DocxPt(size)
    style.font.bold = bold
    style.font.color.rgb = DocxRGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "PingFang SC")


def word_bytes(title: str, sections: list[dict[str, Any]]) -> bytes:
    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = Cm(2.4)
    page.left_margin = page.right_margin = Cm(2.5)
    _font(doc.styles["Normal"], 11)
    _font(doc.styles["Heading 1"], 13, True, "27698E")
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.45
    normal.space_after = DocxPt(8)
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = DocxPt(20)
    title_run = heading.add_run(title)
    title_run.bold = True
    title_run.font.name = "PingFang SC"
    title_run.font.size = DocxPt(18)
    title_run.font.color.rgb = DocxRGBColor.from_string("183B59")
    title_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
    for section in sections:
        if section.get("heading"):
            doc.add_paragraph(section["heading"], style="Heading 1")
        for paragraph in section["paragraphs"]:
            doc.add_paragraph(paragraph)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def draft_sections(content: str) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    for raw in content.splitlines():
        line = raw.strip()
        if not line or line.startswith("# "):
            continue
        if line == "---":
            continue
        if re.match(r"^[一二三四五六七八九十]+、", line):
            sections.append({"heading": line, "paragraphs": []})
        else:
            if not sections:
                sections.append({"heading": "", "paragraphs": []})
            sections[-1]["paragraphs"].append(line)
    return sections


def draft_word_bytes(title: str, content: str) -> bytes:
    return word_bytes(title, draft_sections(content))


def _safe_excel(value: Any) -> str | int | float:
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def excel_bytes(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "资料清单"
    last = get_column_letter(len(columns))
    sheet.merge_cells(f"A1:{last}1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="PingFang SC", size=17, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F597B")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 38
    sheet.merge_cells(f"A2:{last}2")
    sheet["A2"] = "可编辑工作表 · 请核对事实和来源，打印前检查分页。"
    sheet["A2"].font = Font(name="PingFang SC", size=10, color="6B8295")
    sheet.row_dimensions[2].height = 25
    for index, header in enumerate(columns, 1):
        cell = sheet.cell(3, index, header)
        cell.fill = PatternFill("solid", fgColor="DDECF6")
        cell.font = Font(name="PingFang SC", size=10, bold=True, color="244966")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        values = [str(row[index - 1]) for row in rows if len(row) >= index]
        sheet.column_dimensions[get_column_letter(index)].width = min(32, max(16, len(header) * 2 + 4, *(min(len(value) + 2, 32) for value in values)))
    sheet.row_dimensions[3].height = 30
    for row_index, row in enumerate(rows, 4):
        for col_index, value in enumerate(row, 1):
            cell = sheet.cell(row_index, col_index, _safe_excel(value))
            cell.font = Font(name="PingFang SC", size=10, color="304F66")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F9FC")
        max_length = max((len(str(value)) for value in row), default=0)
        sheet.row_dimensions[row_index].height = min(120, max(32, 15 * math.ceil(max_length / 24)))
    end = max(4, 3 + len(rows))
    sheet.auto_filter.ref = f"A3:{last}{end}"
    sheet.freeze_panes = "A4"
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:3"
    sheet.print_area = f"A1:{last}{end}"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _add_text(slide, x, y, w, h, text: str, *, size: int, bold: bool = False, color: tuple[int, int, int] = (28, 61, 86)) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.name = "PingFang SC"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def ppt_bytes(title: str, slides: list[dict[str, Any]]) -> bytes:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]
    for index, item in enumerate([{"title": title, "bullets": ["事实与材料整理", "请核对来源后使用"]}, *slides]):
        slide = deck.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(247, 251, 254)
        bar = slide.shapes.add_shape(1, Inches(.65), Inches(.7), Inches(.09), Inches(.46))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(56, 145, 196)
        bar.line.fill.background()
        _add_text(slide, .88, .68, 11.4, 1.12, item["title"], size=28 if index == 0 else 25, bold=True)
        for bullet_index, bullet in enumerate(item["bullets"][:6]):
            y = (2.05 if index == 0 else 1.82) + bullet_index * .83
            _add_text(slide, .88, y, 11.4, .72, f"•  {bullet}", size=17)
        _add_text(slide, .88, 7.0, 10, .27, "法寻 · 可编辑简报 · 使用前核对事实", size=10, color=(123, 151, 170))
        _add_text(slide, 12.0, 7.0, .5, .27, str(index + 1).zfill(2), size=10, color=(123, 151, 170))
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def template_payload(template_id: str) -> tuple[str, dict[str, Any]]:
    if template_id == "rental-contract":
        _, content = documents.render("rental-contract", {})
        return "word", {"title": "房屋租赁合同（空白模板）", "sections": draft_sections(content)}
    if template_id == "evidence-ledger":
        return "excel", {"title": "案件证据清单", "columns": ["序号", "证据名称", "形成时间", "来源", "待证明事实", "文件位置", "备注"],
                         "rows": [[index, "", "", "", "", "", ""] for index in range(1, 11)]}
    if template_id == "case-brief":
        return "ppt", {"title": "案件沟通简报", "slides": [
            {"title": "事项概览", "bullets": ["当事人及相关主体：待填写", "争议事项：待填写", "希望解决的问题：待填写"]},
            {"title": "事实时间线", "bullets": ["时间与事件 1：待填写", "时间与事件 2：待填写", "现有证据：待填写"]},
            {"title": "待确认问题", "bullets": ["需要核对的事实：待填写", "需要咨询律师的问题：待填写", "下一步行动：待填写"]},
        ]}
    raise KeyError(template_id)


def file_bytes(kind: str, payload: dict[str, Any]) -> bytes:
    if kind == "word":
        return word_bytes(payload["title"], payload["sections"])
    if kind == "excel":
        return excel_bytes(payload["title"], payload["columns"], payload["rows"])
    if kind == "ppt":
        return ppt_bytes(payload["title"], payload["slides"])
    raise ValueError("Unsupported Office type")


def print_html(kind: str, payload: dict[str, Any]) -> str:
    esc = html.escape
    if kind == "word":
        content = "".join((f"<h2>{esc(part['heading'])}</h2>" if part.get("heading") else "") +
                          "".join(f"<p>{esc(text)}</p>" for text in part["paragraphs"]) for part in payload["sections"])
    elif kind == "excel":
        head = "".join(f"<th>{esc(col)}</th>" for col in payload["columns"])
        rows = "".join("<tr>" + "".join(f"<td>{esc(str(cell))}</td>" for cell in row) + "</tr>" for row in payload["rows"])
        content = f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>"
    elif kind == "ppt":
        content = "".join(f"<section class='slide'><h2>{esc(slide['title'])}</h2><ul>" +
                          "".join(f"<li>{esc(item)}</li>" for item in slide["bullets"]) + "</ul></section>" for slide in payload["slides"])
    else:
        raise ValueError("Unsupported Office type")
    return f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>{esc(payload['title'])}</title>
<style>@page{{size:A4;margin:18mm}}body{{font:14px/1.75 'PingFang SC','Microsoft YaHei',sans-serif;color:#213c52;max-width:900px;margin:28px auto;padding:0 18px}}h1{{font-size:25px;text-align:center}}h2{{font-size:18px;margin-top:24px}}p{{white-space:pre-wrap}}table{{border-collapse:collapse;width:100%;font-size:11px}}th,td{{border:1px solid #a9bfd0;padding:6px;text-align:left;vertical-align:top}}th{{background:#e6f1f8}}.slide{{min-height:600px;padding:20px 30px;break-after:page}}.slide li{{font-size:20px;margin:18px 0}}button{{margin-bottom:20px;padding:9px 16px;border:0;border-radius:8px;background:#2d81b9;color:#fff;cursor:pointer}}@media print{{button{{display:none}}body{{margin:0;padding:0}}}}</style></head>
<body><button type="button" onclick="window.print()">打印</button><h1>{esc(payload['title'])}</h1>{content}</body></html>"""
