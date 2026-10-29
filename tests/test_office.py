import io
import os
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from legal_service.app import create_app
from legal_service.documents import render
from legal_service.office_files import excel_bytes


def ai_payload(kind):
    if kind == "word":
        return {"title": "租赁事项说明", "sections": [{"heading": "已知事实", "paragraphs": ["出租人张三与承租人李四拟租赁房屋。"]}]}
    if kind == "excel":
        return {"title": "交接清单", "columns": ["项目", "状态"], "rows": [["钥匙", "待确认"]]}
    return {"title": "事项简报", "slides": [{"title": "背景", "bullets": ["双方拟就租赁事项沟通。"]}]}


class OfficeFlowTests(unittest.TestCase):
    def test_rental_contract_prefill_and_validation(self):
        _, content = render("rental-contract", {"landlord": "张三", "tenant": "李四", "rent": "5000"})
        self.assertIn("出租人（甲方）：张三", content)
        self.assertIn("人民币 5000 元", content)
        self.assertIn("【房屋详细地址待确认】", content)
        with self.assertRaises(Exception):
            render("rental-contract", {"start_date": "2027-01-01", "end_date": "2026-01-01"})
        with self.assertRaises(Exception):
            render("rental-contract", {"deposit": "-1"})

    def test_template_files_private_exports_and_ai_files(self):
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, {"APP_DATA_DIR": folder}):
            with TestClient(create_app()) as app:
                template = app.get("/api/v1/office/templates").json()["items"]
                self.assertEqual({item["kind"] for item in template}, {"word", "excel", "ppt"})
                word = app.get("/api/v1/office/templates/rental-contract/download")
                self.assertEqual(word.status_code, 200)
                self.assertIn("【出租人姓名或名称待确认】", "\n".join(p.text for p in Document(io.BytesIO(word.content)).paragraphs))
                self.assertIn("打印", app.get("/api/v1/office/templates/rental-contract/print").text)
                excel = app.get("/api/v1/office/templates/evidence-ledger/download")
                book = load_workbook(io.BytesIO(excel.content))
                self.assertEqual(book.active["B3"].value, "证据名称")
                ppt = app.get("/api/v1/office/templates/case-brief/download")
                self.assertEqual(len(Presentation(io.BytesIO(ppt.content)).slides), 4)
                self.assertEqual(app.get("/api/v1/office/templates/unknown/download").status_code, 404)

                owner = TestClient(app.app)
                outsider = TestClient(app.app)
                a = owner.post("/api/v1/auth/register", json={"email":"word-owner@example.com","display_name":"文书用户","password":"long-test-password-123"})
                b = outsider.post("/api/v1/auth/register", json={"email":"word-other@example.com","display_name":"其他用户","password":"long-test-password-123"})
                owner_header = {"X-CSRF-Token": a.json()["csrf_token"]}
                other_header = {"X-CSRF-Token": b.json()["csrf_token"]}
                created = owner.post("/api/v1/documents/drafts", json={"template_id":"rental-contract","fields":{"landlord":"张三","tenant":"李四","rent":"5000"}}, headers=owner_header)
                self.assertEqual(created.status_code, 201, created.text)
                draft_id = created.json()["id"]
                draft_doc = owner.get(f"/api/v1/documents/drafts/{draft_id}/download")
                self.assertEqual(draft_doc.status_code, 200)
                self.assertIn("张三", "\n".join(p.text for p in Document(io.BytesIO(draft_doc.content)).paragraphs))
                self.assertEqual(outsider.get(f"/api/v1/documents/drafts/{draft_id}/download").status_code, 404)
                self.assertIn("张三", owner.get(f"/api/v1/documents/drafts/{draft_id}/print").text)
                self.assertEqual(outsider.get(f"/api/v1/documents/drafts/{draft_id}/print").status_code, 404)

                with patch("legal_service.office_ai.generate", new=AsyncMock(side_effect=[ai_payload("word"), ai_payload("excel"), ai_payload("ppt")])):
                    for kind in ("word", "excel", "ppt"):
                        self.assertEqual(owner.post("/api/v1/office/files", json={"kind":kind,"prompt":"整理一份需要双方核对的房屋租赁沟通材料，包含事实和待确认事项。"}).status_code, 403)
                        result = owner.post("/api/v1/office/files", json={"kind":kind,"prompt":"整理一份需要双方核对的房屋租赁沟通材料，包含事实和待确认事项。"}, headers=owner_header)
                        self.assertEqual(result.status_code, 201, result.text)
                        artifact_id = result.json()["id"]
                        file_url = f"/api/v1/office/files/{artifact_id}"
                        self.assertEqual(outsider.get(file_url + "/download").status_code, 404)
                        self.assertEqual(outsider.get(file_url + "/print").status_code, 404)
                        self.assertEqual(owner.get(file_url + "/download").status_code, 200)
                        self.assertIn("打印", owner.get(file_url + "/print").text)
                self.assertEqual(len(owner.get("/api/v1/office/files").json()["items"]), 3)
                self.assertEqual(outsider.get("/api/v1/office/files").json()["items"], [])

    def test_excel_formula_like_text_is_not_a_formula(self):
        content = excel_bytes("安全清单", ["项目", "说明"], [["=1+1", "@外部输入"]])
        sheet = load_workbook(io.BytesIO(content)).active
        self.assertEqual(sheet["A4"].data_type, "s")
        self.assertTrue(sheet["A4"].value.startswith("'"))


if __name__ == "__main__":
    unittest.main()
