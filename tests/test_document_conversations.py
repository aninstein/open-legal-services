import io
import os
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from docx import Document
from fastapi.testclient import TestClient

from legal_service.app import create_app
from legal_service.document_scenarios import CATALOG, render_scenario


class DocumentConversationTests(unittest.TestCase):
    def test_catalog_blueprints_have_intake_and_renderable_placeholders(self):
        self.assertGreaterEqual(len(CATALOG), 40)
        for item in CATALOG.values():
            self.assertIn('必填信息', item['intake_prompt'])
            self.assertIn('非常建议提供', item['intake_prompt'])
            self.assertIn('选填信息', item['intake_prompt'])
            facts = {field['key']: '已提供的事实' for field in item['fields'] if field['level'] == 'required'}
            _, content = render_scenario(item['id'], facts)
            self.assertIn('已提供的事实', content)
            self.assertNotIn('{', content, item['id'])

    def test_chat_generates_excel_and_ppt_files(self):
        from openpyxl import load_workbook
        from pptx import Presentation
        cases = [
            ('evidence-ledger-chat', {'matter':'租赁纠纷','items':'租赁合同；转账截图','proof':'双方租赁关系和付款'}, 'excel'),
            ('case-brief-chat', {'matter':'租赁纠纷','parties':'张三、李四','facts':'双方就押金返还发生争议'}, 'ppt'),
        ]
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, {'APP_DATA_DIR': folder}):
            with TestClient(create_app()) as app:
                owner = TestClient(app.app)
                a = owner.post('/api/v1/auth/register', json={'email':'office-chat@example.com','display_name':'文书用户','password':'long-test-password-123'})
                header = {'X-CSRF-Token':a.json()['csrf_token']}
                for scenario_id, facts, kind in cases:
                    conv = owner.post('/api/v1/documents/conversations',json={'scenario_id':scenario_id},headers=header).json()
                    with patch('legal_service.document_intake.extract', new=AsyncMock(return_value=facts)):
                        response = owner.post(f"/api/v1/documents/conversations/{conv['id']}/messages",json={'body':'我的情况都在这里'},headers=header)
                    self.assertEqual(response.status_code, 200, response.text)
                    item = response.json()
                    self.assertEqual(item['kind'], kind)
                    file_response = owner.get(f"/api/v1/office/files/{item['artifact_id']}/download")
                    self.assertEqual(file_response.status_code, 200)
                    if kind == 'excel':
                        self.assertEqual(load_workbook(io.BytesIO(file_response.content)).active['B4'].value, '租赁合同')
                    else:
                        self.assertGreaterEqual(len(Presentation(io.BytesIO(file_response.content)).slides), 4)

    def test_conversation_asks_missing_then_builds_private_word(self):
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, {'APP_DATA_DIR': folder}):
            with TestClient(create_app()) as app:
                owner, outsider = TestClient(app.app), TestClient(app.app)
                a = owner.post('/api/v1/auth/register', json={'email':'doc-owner@example.com','display_name':'文书用户','password':'long-test-password-123'})
                b = outsider.post('/api/v1/auth/register', json={'email':'doc-other@example.com','display_name':'其他用户','password':'long-test-password-123'})
                header = {'X-CSRF-Token':a.json()['csrf_token']}
                other = {'X-CSRF-Token':b.json()['csrf_token']}
                catalog = app.get('/api/v1/documents/scenarios?q=租赁').json()['items']
                self.assertTrue(any(item['id'] == 'rental-contract' for item in catalog))
                self.assertIn('intake_prompt', app.get('/api/v1/documents/scenarios/rental-contract').json())
                self.assertEqual(owner.post('/api/v1/documents/conversations',json={'scenario_id':'rental-contract'}).status_code, 403)
                started = owner.post('/api/v1/documents/conversations',json={'scenario_id':'rental-contract'},headers=header)
                self.assertEqual(started.status_code, 201, started.text)
                conv_id = started.json()['id']
                self.assertIn('出租人', started.json()['messages'][0]['content'])
                self.assertEqual(outsider.get(f'/api/v1/documents/conversations/{conv_id}').status_code, 404)
                with patch('legal_service.document_intake.extract', new=AsyncMock(side_effect=[
                    {'landlord':'张三','tenant':'李四'},
                    {'property':'北京市朝阳区示例路1号','term':'2027年1月至12月','rent':'每月5000元','deposit':'5000元，交还后7日返还'}
                ])):
                    first = owner.post(f'/api/v1/documents/conversations/{conv_id}/messages',json={'body':'房东张三，租客李四'},headers=header)
                    self.assertEqual(first.status_code, 200, first.text)
                    self.assertIsNone(first.json()['draft_id'])
                    self.assertIn('房屋地址', first.json()['messages'][-1]['content'])
                    second = owner.post(f'/api/v1/documents/conversations/{conv_id}/messages',json={'body':'房子在北京市朝阳区示例路1号，租期2027年1月至12月，月租5000，押金5000元，交还后7日返还'},headers=header)
                    self.assertEqual(second.status_code, 200, second.text)
                draft_id = second.json()['draft_id']
                self.assertTrue(draft_id)
                self.assertEqual(outsider.get(f'/api/v1/documents/drafts/{draft_id}/download').status_code, 404)
                content = owner.get(f'/api/v1/documents/drafts/{draft_id}/download').content
                paragraphs = '\n'.join(p.text for p in Document(io.BytesIO(content)).paragraphs)
                self.assertIn('张三', paragraphs)
                self.assertIn('北京市朝阳区示例路1号', paragraphs)
                self.assertIn('交还后7日返还', paragraphs)
                self.assertEqual(owner.get(f'/api/v1/documents/conversations/{conv_id}').json()['draft_id'], draft_id)
                self.assertEqual(len(owner.get('/api/v1/documents/conversations').json()['items']), 1)


if __name__ == '__main__':
    unittest.main()
